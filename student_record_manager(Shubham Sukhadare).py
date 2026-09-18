"""
Student Record Management System
---------------------------------
A menu-driven console application to Add, View, Search, Update,
and Delete student records. Records are stored permanently in a
JSON file (students.json) so they survive between program runs.

Concepts demonstrated:
- Data types & variables
- Conditional statements (if/elif/else)
- Loops (while, for)
- Functions
- Exception handling (try/except)
- File I/O (json module)
- Menu-driven console design
"""

import json
import os

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

FILE_NAME = "students.json"

# ---------------------------------------------------------
# SIMPLE COLOR / STYLE HELPERS (using ANSI escape codes)
# These just make text appear in color in most terminals.
# No external library needed.
# ---------------------------------------------------------

class Color:
    RESET = "\033[0m"
    BOLD = "\033[1m"
    CYAN = "\033[96m"
    GREEN = "\033[92m"
    YELLOW = "\033[93m"
    RED = "\033[91m"
    MAGENTA = "\033[95m"
    BLUE = "\033[94m"


def success(message):
    print(f"{Color.GREEN}✔ {message}{Color.RESET}")


def error(message):
    print(f"{Color.RED}✘ {message}{Color.RESET}")


def warning(message):
    print(f"{Color.YELLOW}⚠ {message}{Color.RESET}")


def info(message):
    print(f"{Color.CYAN}{message}{Color.RESET}")


def print_banner(text, color=Color.CYAN, width=46):
    """Print text centered inside a box made of Unicode box-drawing characters."""
    print(color + "╔" + "═" * width + "╗")
    print("║" + text.center(width) + "║")
    print("╚" + "═" * width + "╝" + Color.RESET)


def print_divider(char="─", width=48, color=Color.BLUE):
    print(color + char * width + Color.RESET)


# ---------------------------------------------------------
# FILE HANDLING FUNCTIONS
# ---------------------------------------------------------

def load_records():
    """
    Load student records from the JSON file into a list of dictionaries.
    If the file doesn't exist yet, or contains invalid/corrupted data,
    start with an empty list instead of crashing.
    """
    if not os.path.exists(FILE_NAME):
        # First time running the program - no file yet, that's fine.
        return []

    try:
        with open(FILE_NAME, "r") as file:
            data = json.load(file)
            if isinstance(data, list):
                return data
            else:
                warning("students.json has an unexpected format. Starting fresh.")
                return []
    except json.JSONDecodeError:
        warning("students.json is corrupted or empty. Starting fresh.")
        return []
    except IOError:
        warning("Could not read students.json. Starting fresh.")
        return []


def save_records(students):
    """
    Save the current list of student records to the JSON file.
    Called after every Add, Update, or Delete operation.
    """
    try:
        with open(FILE_NAME, "w") as file:
            json.dump(students, file, indent=4)
    except IOError:
        error("Could not save data to file. Your changes may be lost.")


# ---------------------------------------------------------
# VALIDATION HELPER FUNCTIONS
# ---------------------------------------------------------

def get_non_empty_input(prompt):
    """Keep asking until the user types something (not just spaces)."""
    while True:
        value = input(f"{Color.CYAN}{prompt}{Color.RESET}").strip()
        if value == "":
            warning("This field cannot be empty. Please try again.")
        else:
            return value


def get_valid_age():
    """Keep asking until the user enters a valid whole number for age."""
    while True:
        age_input = input(f"{Color.CYAN}Enter Age: {Color.RESET}").strip()
        try:
            age = int(age_input)
            if age <= 0 or age > 100:
                warning("Please enter a realistic age (1-100).")
                continue
            return age
        except ValueError:
            warning("Invalid input. Age must be a number (e.g., 22). Try again.")


def find_student_by_id(students, student_id):
    """
    Search the list for a student with the given ID.
    Returns the matching dictionary, or None if not found.
    """
    for student in students:
        if student["id"] == student_id:
            return student
    return None


# ---------------------------------------------------------
# DISPLAY HELPER
# ---------------------------------------------------------

def print_student_card(student):
    print_divider("─", 48, Color.MAGENTA)
    print(f"{Color.BOLD} 🎓 ID     :{Color.RESET} {student['id']}")
    print(f"{Color.BOLD} 👤 Name   :{Color.RESET} {student['name']}")
    print(f"{Color.BOLD} 🎂 Age    :{Color.RESET} {student['age']}")
    print(f"{Color.BOLD} 📘 Course :{Color.RESET} {student['course']}")
    print(f"{Color.BOLD} ✉️  Email  :{Color.RESET} {student['email']}")
    print(f"{Color.BOLD} 📞 Phone  :{Color.RESET} {student['phone']}")
    print_divider("─", 48, Color.MAGENTA)


# ---------------------------------------------------------
# CORE FEATURE FUNCTIONS
# ---------------------------------------------------------

def add_record(students):
    print_banner(" ADD NEW STUDENT ", Color.GREEN)
    student_id = get_non_empty_input("Enter Student ID: ")

    if find_student_by_id(students, student_id) is not None:
        error(f"A student with ID '{student_id}' already exists.")
        return

    name = get_non_empty_input("Enter Name: ")
    age = get_valid_age()
    course = get_non_empty_input("Enter Course: ")
    email = get_non_empty_input("Enter Email: ")
    phone = get_non_empty_input("Enter Phone Number: ")

    new_student = {
        "id": student_id,
        "name": name,
        "age": age,
        "course": course,
        "email": email,
        "phone": phone,
    }

    students.append(new_student)
    save_records(students)
    success(f"Student '{name}' added successfully!")


def view_records(students):
    print_banner(" ALL STUDENT RECORDS ", Color.BLUE)
    if not students:
        warning("No records found. Add one from the main menu!")
        return

    print(f"{Color.BOLD}Total students: {len(students)}{Color.RESET}")
    for student in students:
        print_student_card(student)


def search_record(students):
    print_banner(" SEARCH STUDENT ", Color.CYAN)
    student_id = get_non_empty_input("Enter Student ID to search: ")
    student = find_student_by_id(students, student_id)

    if student is None:
        error(f"No student found with ID '{student_id}'.")
        return

    success("Student found!")
    print_student_card(student)


def update_record(students):
    print_banner(" UPDATE STUDENT ", Color.YELLOW)
    student_id = get_non_empty_input("Enter Student ID to update: ")
    student = find_student_by_id(students, student_id)

    if student is None:
        error(f"No student found with ID '{student_id}'.")
        return

    info("Leave a field blank to keep its current value.")

    new_name = input(f"{Color.CYAN}Enter new Name [{student['name']}]: {Color.RESET}").strip()
    if new_name != "":
        student["name"] = new_name

    new_age = input(f"{Color.CYAN}Enter new Age [{student['age']}]: {Color.RESET}").strip()
    if new_age != "":
        try:
            student["age"] = int(new_age)
        except ValueError:
            warning("Invalid age entered. Keeping the old age.")

    new_course = input(f"{Color.CYAN}Enter new Course [{student['course']}]: {Color.RESET}").strip()
    if new_course != "":
        student["course"] = new_course

    new_email = input(f"{Color.CYAN}Enter new Email [{student['email']}]: {Color.RESET}").strip()
    if new_email != "":
        student["email"] = new_email

    new_phone = input(f"{Color.CYAN}Enter new Phone [{student['phone']}]: {Color.RESET}").strip()
    if new_phone != "":
        student["phone"] = new_phone

    save_records(students)
    success("Record updated successfully!")


def delete_record(students):
    print_banner(" DELETE STUDENT ", Color.RED)
    student_id = get_non_empty_input("Enter Student ID to delete: ")
    student = find_student_by_id(students, student_id)

    if student is None:
        error(f"No student found with ID '{student_id}'.")
        return

    print_student_card(student)
    confirm = input(f"{Color.YELLOW}Are you sure you want to delete this student? (y/n): {Color.RESET}").strip().lower()
    if confirm == "y":
        students.remove(student)
        save_records(students)
        success("Record deleted successfully!")
    else:
        info("Delete cancelled.")


def export_to_excel(students):
    """
    Write all student records into a real Excel file (students_report.xlsx)
    using the openpyxl library. This is separate from students.json
    (which is the 'real' storage used to reload data) - the .xlsx file
    is just a formatted report you can open directly in Excel.
    """
    print_banner(" EXPORT RECORDS TO EXCEL ", Color.BLUE)

    if not OPENPYXL_AVAILABLE:
        error("The 'openpyxl' library is required for Excel export but isn't installed.")
        info("Install it with:  pip install openpyxl")
        return

    if not students:
        warning("No records to export. Add some students first.")
        return

    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Students"

        headers = ["ID", "Name", "Age", "Course", "Email", "Phone"]
        sheet.append(headers)

        # Style the header row: bold white text on a dark blue fill.
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # One row per student record.
        for student in students:
            sheet.append([
                student["id"],
                student["name"],
                student["age"],
                student["course"],
                student["email"],
                student["phone"],
            ])

        # Auto-fit column widths based on the longest value in each column.
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value)) for cell in column_cells)
            column_letter = column_cells[0].column_letter
            sheet.column_dimensions[column_letter].width = max_length + 4

        workbook.save("students_report.xlsx")
        success(f"Exported {len(students)} record(s) to students_report.xlsx")
    except IOError:
        error("Could not write to students_report.xlsx. Is the file open in Excel?")
    except Exception as e:
        error(f"Unexpected error during export: {e}")


# ---------------------------------------------------------
# MENU / MAIN PROGRAM
# ---------------------------------------------------------

def show_menu():
    print()
    print_banner(" 🎓 STUDENT RECORD MANAGEMENT SYSTEM 🎓 ", Color.MAGENTA, width=50)
    print(f"""{Color.BOLD}
  1. ➕  Add Record
  2. 📋  View Records
  3. 🔍  Search Record
  4. ✏️   Update Record
  5. 🗑️   Delete Record
  6. 📊  Export Records to Excel
  7. 🚪  Exit
{Color.RESET}""")
    print_divider("═", 52, Color.MAGENTA)


def main():
    students = load_records()

    print_banner(" WELCOME ", Color.GREEN, width=50)
    info(f"Loaded {len(students)} existing record(s) from {FILE_NAME}.")

    while True:
        show_menu()
        choice = input(f"{Color.BOLD}Enter your choice (1-7): {Color.RESET}").strip()

        try:
            choice = int(choice)
        except ValueError:
            error("Invalid input. Please enter a number between 1 and 7.")
            continue

        if choice == 1:
            add_record(students)
        elif choice == 2:
            view_records(students)
        elif choice == 3:
            search_record(students)
        elif choice == 4:
            update_record(students)
        elif choice == 5:
            delete_record(students)
        elif choice == 6:
            export_to_excel(students)
        elif choice == 7:
            print_banner(" GOODBYE! 👋 ", Color.CYAN, width=50)
            break
        else:
            error("Invalid choice. Please select a number between 1 and 7.")


if __name__ == "__main__":
    main()
